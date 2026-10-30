
import json
from datetime import datetime, date
from pathlib import Path

from openpyxl import load_workbook


# Пути к файлам проекта
BASE_DIR = Path(__file__).resolve().parent
EXCEL_FILE = BASE_DIR / "clinical_recommendations.xlsx"
OUTPUT_FILE = BASE_DIR / "catalog.json"

# Название листа Excel
SHEET_NAME = "Лист1"

# Ожидаемые заголовки столбцов
EXPECTED_COLUMNS = [
    "ID",
    "Наименование",
    "МКБ-10",
    "Возрастная категория",
    "Разработчик",
    "Статус одобрения НПС",
    "Дата размещения",
    "Статус применения",
]


def clean_value(value):
    """Преобразует значение Excel в текст."""
    if value is None:
        return ""

    return str(value).strip()


def format_date(value):
    """Преобразует дату Excel в единый формат."""
    if value is None:
        return None

    if isinstance(value, (datetime, date)):
        return value.isoformat()

    return str(value).strip()


def main():
    if not EXCEL_FILE.exists():
        raise FileNotFoundError(
            f"Не найден файл Excel: {EXCEL_FILE.name}"
        )

    print(f"Читаем файл: {EXCEL_FILE.name}")

    workbook = load_workbook(
        EXCEL_FILE,
        read_only=True,
        data_only=True
    )

    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(
            f"Лист '{SHEET_NAME}' не найден. "
            f"Доступные листы: {workbook.sheetnames}"
        )

    sheet = workbook[SHEET_NAME]

    rows = sheet.iter_rows(values_only=True)

    headers = next(rows, None)

    if headers is None:
        raise ValueError("Excel-файл пустой.")

    actual_columns = [
        clean_value(value) for value in headers
    ]

    if actual_columns != EXPECTED_COLUMNS:
        raise ValueError(
            "Структура Excel отличается от ожидаемой.\n"
            f"Ожидаемые столбцы: {EXPECTED_COLUMNS}\n"
            f"Полученные столбцы: {actual_columns}"
        )

    recommendations = []
    errors = []

    for row_number, row in enumerate(rows, start=2):

        if not row or not any(value is not None for value in row):
            continue

        if len(row) < len(EXPECTED_COLUMNS):
            errors.append(
                f"Строка {row_number}: недостаточно столбцов."
            )
            continue

        recommendation_id = clean_value(row[0])
        title = clean_value(row[1])

        if not recommendation_id or not title:
            errors.append(
                f"Строка {row_number}: отсутствует ID или название."
            )
            continue

        recommendation = {
            "id": recommendation_id,
            "title": title,
            "icd10": clean_value(row[2]),
            "age_category": clean_value(row[3]),
            "developer": clean_value(row[4]),
            "nps_approved": clean_value(row[5]),
            "publication_date": format_date(row[6]),
            "application_status": clean_value(row[7]),
        }

        recommendations.append(recommendation)

    workbook.close()

    if errors:
        print("Обнаружены ошибки:")
        for error in errors:
            print(error)

        raise ValueError(
            f"Не удалось обработать {len(errors)} строк."
        )

    # Проверяем уникальность ID
    ids = [
        item["id"] for item in recommendations
    ]

    if len(ids) != len(set(ids)):
        raise ValueError(
            "В Excel обнаружены повторяющиеся ID рекомендаций."
        )

    catalog = {
        "updated_at": datetime.now().astimezone().isoformat(),
        "source": "https://cr.minzdrav.gov.ru/clin-rec",
        "source_file": EXCEL_FILE.name,
        "total": len(recommendations),
        "recommendations": recommendations,
    }

    with open(OUTPUT_FILE, "w", encoding="utf-8") as file:
        json.dump(
            catalog,
            file,
            ensure_ascii=False,
            indent=2
        )

    print("--------------------------------")
    print("Каталог успешно сформирован.")
    print(f"Количество рекомендаций: {len(recommendations)}")
    print(f"Файл результата: {OUTPUT_FILE.name}")
    print("--------------------------------")


if __name__ == "__main__":
    main()
